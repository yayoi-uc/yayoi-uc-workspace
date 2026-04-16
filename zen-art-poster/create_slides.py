#!/usr/bin/env python3
"""zen art poster - Slide Deck Generator"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# === Colors ===
KINARI = RGBColor(0xF5, 0xF0, 0xE8)
WARM_WHITE = RGBColor(0xFA, 0xFA, 0xF7)
SUMI = RGBColor(0x2C, 0x2C, 0x2C)
WARM_GRAY = RGBColor(0x6B, 0x65, 0x60)
LIGHT_GRAY = RGBColor(0xB5, 0xB0, 0xA8)
AI_BLUE = RGBColor(0x2B, 0x4C, 0x7E)
SAND = RGBColor(0xC8, 0xBF, 0xA9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# Use blank layout
blank_layout = prs.slide_layouts[6]


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.5)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = SAND
    shape.line.fill.background()
    return shape


def add_text(slide, left, top, width, height, text, font_name='Helvetica Neue',
             font_size=14, color=SUMI, bold=False, alignment=PP_ALIGN.LEFT,
             font_name_jp=None, anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    txBox.text_frame.word_wrap = True
    txBox.text_frame.auto_size = None
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.font.name = font_name
    if font_name_jp:
        p.font.name = font_name_jp
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return txBox


def add_multiline(slide, left, top, width, height, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, font_name, font_size, color, bold, alignment)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for i, (text, fname, fsize, color, bold, align) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.name = fname
        p.font.size = Pt(fsize)
        p.font.color.rgb = color
        p.font.bold = bold
        p.alignment = align
        p.space_after = Pt(2)
    return txBox


def add_photo_placeholder(slide, left, top, width, height, label="PHOTO"):
    shape = add_rect(slide, left, top, width, height, RGBColor(0xE8, 0xE4, 0xDC))
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.name = 'Helvetica Neue'
    p.font.size = Pt(11)
    p.font.color.rgb = LIGHT_GRAY
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    shape.text_frame.auto_size = None
    return shape


# Helpers
L_MARGIN = Inches(1.2)
R_HALF = Inches(6.666)
HALF_W = Inches(6.666)
FULL_W = Inches(10.9)
CAPTION_SIZE = 11
H1_SIZE = 42
H2_SIZE = 34
H3_SIZE = 24
SUB_EN_SIZE = 18
SUB_JP_SIZE = 15
BODY_EN_SIZE = 15
BODY_JP_SIZE = 13


# ================================================================
# SLIDE 1: Cover
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

# Right photo placeholder
add_photo_placeholder(slide, R_HALF, Inches(0), HALF_W, SLIDE_H,
                      "PHOTO\nkakejiku in minimal interior")

# Left text
add_text(slide, L_MARGIN, Inches(2.0), Inches(5), Inches(0.4),
         "ZEN ART POSTER", font_size=CAPTION_SIZE, color=LIGHT_GRAY,
         alignment=PP_ALIGN.CENTER)

add_text(slide, L_MARGIN, Inches(2.6), Inches(5), Inches(1.2),
         "A Portable Window\nof Zen", font_size=H1_SIZE, color=SUMI,
         alignment=PP_ALIGN.CENTER)

add_line(slide, Inches(3.0), Inches(4.0), Inches(0.6))

add_text(slide, L_MARGIN, Inches(4.3), Inches(5), Inches(0.5),
         "禅を持ち運ぶ、小さな窓", font_size=SUB_JP_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)


# ================================================================
# SLIDE 2: Opening Story
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, SUMI)

add_text(slide, Inches(1.5), Inches(2.4), Inches(10.3), Inches(1.0),
         '"I changed the scroll for New Year\'s,"\nmy mother said.',
         font_size=22, color=KINARI, alignment=PP_ALIGN.CENTER)

add_line(slide, Inches(6.2), Inches(3.8), Inches(0.9))
# override line color for dark bg
for shape in slide.shapes:
    if shape.shape_type == 1 and shape.height == Pt(1):
        shape.fill.fore_color.rgb = WARM_GRAY

add_text(slide, Inches(1.5), Inches(4.2), Inches(10.3), Inches(1.0),
         '「正月用に掛け軸変えたよ」\n母がさりげなく言った。',
         font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# ================================================================
# SLIDE 3: The Realization
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_photo_placeholder(slide, Inches(0), Inches(0), HALF_W, SLIDE_H,
                      "PHOTO\ntokonoma / tea cabinet")

add_text(slide, Inches(7.2), Inches(1.5), Inches(5.2), Inches(1.2),
         "The scroll had always been there.\nBut I never noticed it changed\nwith each season.",
         font_size=SUB_EN_SIZE, color=WARM_GRAY)

add_text(slide, Inches(7.2), Inches(2.9), Inches(5.2), Inches(1.0),
         "床の間の掛け軸はいつもそこにあった。\nけれど季節ごとに変わっていたことには\n気づかなかった。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

add_line(slide, Inches(7.2), Inches(4.1), Inches(0.6))

add_text(slide, Inches(7.2), Inches(4.4), Inches(5.2), Inches(0.8),
         "Portable. Storable. Interchangeable.\nPerhaps the most minimal art in the world.",
         font_size=18, color=SUMI)

add_text(slide, Inches(7.2), Inches(5.4), Inches(5.2), Inches(0.6),
         "持ち運べて、収納できて、交換できる。\n世界でもっともミニマルなアート。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)


# ================================================================
# SLIDE 4: What is Kakejiku?
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, L_MARGIN, Inches(0.8), FULL_W, Inches(0.7),
         "The World's Only Interchangeable Art",
         font_size=H2_SIZE, color=SUMI)

add_text(slide, L_MARGIN, Inches(1.5), FULL_W, Inches(0.4),
         "世界で唯一の「交換型アート」",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

# Table
table_data = [
    ["Western Art", "Kakejiku"],
    ["Fixed on the wall", "Changed with the seasons"],
    ["Heavy frame", "Rolls up, stores in a box"],
    ["One piece, one space", "Many pieces, one space"],
    ["Decorative", "Intentional — chosen for the moment"],
]

tbl = slide.shapes.add_table(len(table_data), 2, L_MARGIN, Inches(2.3), Inches(9), Inches(2.8)).table
tbl.columns[0].width = Inches(4.5)
tbl.columns[1].width = Inches(4.5)

for i, row_data in enumerate(table_data):
    for j, cell_text in enumerate(row_data):
        cell = tbl.cell(i, j)
        cell.text = cell_text
        p = cell.text_frame.paragraphs[0]
        if i == 0:
            p.font.size = Pt(12)
            p.font.color.rgb = LIGHT_GRAY
            p.font.bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = KINARI
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = SUMI
            p.font.bold = False
            cell.fill.solid()
            cell.fill.fore_color.rgb = KINARI
        p.font.name = 'Helvetica Neue'

add_line(slide, L_MARGIN, Inches(5.5), Inches(0.6))

add_text(slide, L_MARGIN, Inches(5.8), FULL_W, Inches(0.5),
         "Born from tea ceremony and Zen culture.\nNow reimagined for contemporary living.",
         font_size=SUB_EN_SIZE, color=WARM_GRAY)

add_text(slide, L_MARGIN, Inches(6.5), FULL_W, Inches(0.4),
         "茶道と禅の文化から生まれた。いま、現代の暮らしのために再定義する。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)


# ================================================================
# SLIDE 5: Concept
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_photo_placeholder(slide, R_HALF, Inches(0), HALF_W, SLIDE_H,
                      "PHOTO\nseasonal kakejiku in living room")

add_text(slide, L_MARGIN, Inches(1.2), Inches(5), Inches(0.3),
         "CORE CONCEPT", font_size=CAPTION_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(1.6), Inches(5), Inches(0.7),
         "Bring Seasons Home", font_size=28, color=SUMI)

add_text(slide, L_MARGIN, Inches(2.3), Inches(5), Inches(0.4),
         "家に四季を持ち込む", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

add_line(slide, L_MARGIN, Inches(2.9), Inches(0.6))

add_text(slide, L_MARGIN, Inches(3.2), Inches(5.2), Inches(1.0),
         "Through the traditional Japanese format of kakejiku\nand the works of artists around the world,\nwe bring seasons and stillness into the home.",
         font_size=BODY_EN_SIZE, color=WARM_GRAY)

add_text(slide, L_MARGIN, Inches(4.3), Inches(5.2), Inches(0.8),
         "掛け軸という日本の伝統フォーマットを使い、\n世界のアーティストの作品を通して\n「四季」と「静けさ」を家の中に持ち込む。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(5.3), Inches(5.2), Inches(0.6),
         "A kakejiku is not just art —\nit is a small window to feel nature and time.",
         font_size=BODY_EN_SIZE, color=WARM_GRAY)

add_text(slide, L_MARGIN, Inches(6.0), Inches(5.2), Inches(0.5),
         "掛け軸は単なるアートではなく、\n自然や時間を感じるための小さな窓。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)


# ================================================================
# SLIDE 6: Worldview
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, SUMI)

add_text(slide, L_MARGIN, Inches(0.8), FULL_W, Inches(0.3),
         "WORLDVIEW", font_size=CAPTION_SIZE, color=WARM_GRAY)

add_text(slide, L_MARGIN, Inches(1.3), FULL_W, Inches(1.2),
         "Japanese Format\n\u00d7\nGlobal Artists",
         font_size=28, color=KINARI, alignment=PP_ALIGN.CENTER)

add_line(slide, Inches(6.0), Inches(2.8), Inches(1.3))
for shape in slide.shapes:
    if hasattr(shape, 'fill') and shape.height == Pt(1):
        try:
            shape.fill.fore_color.rgb = WARM_GRAY
        except:
            pass

# Left column
add_text(slide, L_MARGIN, Inches(3.3), Inches(5.2), Inches(1.2),
         "Modern living spaces are minimal and global,\nyet the culture of feeling seasons at home\nhas almost disappeared.",
         font_size=BODY_EN_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(4.5), Inches(5.2), Inches(0.8),
         "現代の住空間はミニマルでグローバル化している一方、\n家の中で季節を感じる文化はほとんど消えている。",
         font_size=BODY_JP_SIZE, color=WARM_GRAY)

# Right column
add_text(slide, Inches(7.2), Inches(3.3), Inches(5.2), Inches(1.2),
         "Japan has a culture of changing art with the seasons.\nWe open this traditional format to artists worldwide\n— building a new art category.",
         font_size=BODY_EN_SIZE, color=LIGHT_GRAY)

add_text(slide, Inches(7.2), Inches(4.5), Inches(5.2), Inches(0.8),
         "日本には「季節ごとにアートを変える文化」がある。\nこの伝統的フォーマットを世界のアーティストに開放し、\n新しいアートカテゴリーとして展開する。",
         font_size=BODY_JP_SIZE, color=WARM_GRAY)

# Spotify line
add_text(slide, Inches(1.5), Inches(5.8), Inches(10.3), Inches(0.5),
         "Like Spotify for your walls — change art with the seasons.",
         font_size=16, color=KINARI, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.4),
         "Spotifyのように、空間のアートを季節で変える体験。",
         font_size=BODY_JP_SIZE, color=WARM_GRAY, alignment=PP_ALIGN.CENTER)


# ================================================================
# SLIDE 7: What We Do
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, L_MARGIN, Inches(0.8), FULL_W, Inches(0.6),
         "What We Do", font_size=H2_SIZE, color=SUMI)

add_text(slide, L_MARGIN, Inches(1.45), FULL_W, Inches(0.4),
         "zen art poster が提案すること", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

# Three cards
card_w = Inches(3.4)
card_h = Inches(3.8)
card_y = Inches(2.3)
gap = Inches(0.35)

cards = [
    ("01", "Select", "Curated vintage kakejiku\nfor the Western interior",
     "欧米インテリアに合う\nヴィンテージ掛け軸のキュレーション"),
    ("02", "Artist Collab", "Global artists paint on\nthe kakejiku format",
     "世界のアーティストが\n掛け軸フォーマットに描く"),
    ("03", "Original", "New canvases from upcycled\nJapanese craft materials",
     "アップサイクル素材による\n新しいキャンバスの開発"),
]

for i, (num, title, desc_en, desc_jp) in enumerate(cards):
    x = L_MARGIN + (card_w + gap) * i
    add_rect(slide, x, card_y, card_w, card_h, WARM_WHITE)
    add_text(slide, x + Inches(0.35), card_y + Inches(0.3), Inches(2.5), Inches(0.3),
             num, font_size=CAPTION_SIZE, color=LIGHT_GRAY)
    add_text(slide, x + Inches(0.35), card_y + Inches(0.7), Inches(2.7), Inches(0.5),
             title, font_size=18, color=SUMI, bold=False)
    add_text(slide, x + Inches(0.35), card_y + Inches(1.3), Inches(2.7), Inches(1.0),
             desc_en, font_size=BODY_EN_SIZE, color=WARM_GRAY)
    add_text(slide, x + Inches(0.35), card_y + Inches(2.4), Inches(2.7), Inches(0.9),
             desc_jp, font_size=BODY_JP_SIZE, color=LIGHT_GRAY)


# ================================================================
# SLIDE 8: Rebone Series
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_photo_placeholder(slide, R_HALF, Inches(0), HALF_W, SLIDE_H,
                      "PHOTO\nmaterial flat-lay / swatches")

add_text(slide, L_MARGIN, Inches(1.2), Inches(5), Inches(0.3),
         "REBONE SERIES", font_size=CAPTION_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(1.7), Inches(5.2), Inches(0.8),
         "Giving Forgotten Materials\na Second Life",
         font_size=28, color=SUMI)

add_text(slide, L_MARGIN, Inches(2.7), Inches(5), Inches(0.4),
         "忘れられた素材に、もう一度命を。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

add_line(slide, L_MARGIN, Inches(3.3), Inches(0.6))

# Pill tags as simple text
pills = "Okayama Denim  |  Kimono Hagire  |  Kokura-ori\nHandmade Washi  |  PVC  |  Tracing Paper"
add_text(slide, L_MARGIN, Inches(3.7), Inches(5.2), Inches(0.8),
         pills, font_size=14, color=WARM_GRAY)

add_text(slide, L_MARGIN, Inches(4.8), Inches(5.2), Inches(0.5),
         "岡山デニム端材 / 着物はぎれ / 小倉織端切れ / 手漉き和紙 / PVC / トレーシングペーパー",
         font_size=11, color=LIGHT_GRAY)


# ================================================================
# SLIDE 9: Washi
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_photo_placeholder(slide, Inches(0), Inches(0), HALF_W, SLIDE_H,
                      "PHOTO\nwashi texture / making process")

add_text(slide, Inches(7.2), Inches(0.8), Inches(5.2), Inches(0.3),
         "MATERIAL", font_size=CAPTION_SIZE, color=LIGHT_GRAY)

add_text(slide, Inches(7.2), Inches(1.2), Inches(5.2), Inches(0.8),
         "Washi: A Canvas That\nLasts Centuries",
         font_size=26, color=SUMI)

add_text(slide, Inches(7.2), Inches(2.1), Inches(5.2), Inches(0.3),
         "何百年も生き続けるキャンバス",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

features = [
    ("Long fibers create a unique bleeding effect", "長い繊維が生む滲みはアーティストに愛される表現"),
    ("No chemical dyes — colors never fade", "化学染料不使用 — 色褪せしない"),
    ("Each region produces a distinct texture", "産地ごとに異なる風合いと個性"),
    ("Naturally sustainable — kozo, mitsumata, gampi", "天然由来の持続可能な素材"),
]

y = 2.8
for en, jp in features:
    add_text(slide, Inches(7.2), Inches(y), Inches(5.2), Inches(0.3),
             en, font_size=BODY_EN_SIZE, color=WARM_GRAY)
    add_text(slide, Inches(7.2), Inches(y + 0.3), Inches(5.2), Inches(0.3),
             jp, font_size=BODY_JP_SIZE, color=LIGHT_GRAY)
    y += 0.75

add_line(slide, Inches(7.2), Inches(y + 0.1), Inches(0.6))

add_text(slide, Inches(7.2), Inches(y + 0.35), Inches(5.2), Inches(0.3),
         "EUDR COMPLIANT — FULL TRACEABILITY",
         font_size=CAPTION_SIZE, color=LIGHT_GRAY)


# ================================================================
# SLIDE 10: Original Kakejiku Concepts
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, L_MARGIN, Inches(0.6), FULL_W, Inches(0.3),
         "ORIGINAL KAKEJIKU", font_size=CAPTION_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(1.0), FULL_W, Inches(0.6),
         "A Canvas Artists Want to Create On",
         font_size=26, color=SUMI)

add_text(slide, L_MARGIN, Inches(1.6), FULL_W, Inches(0.3),
         "アーティストが描きたいと思うキャンバスの提案",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

concepts = [
    ("CONCEPT 01", "Seasonal Rotation",
     "A curated set of scrolls designed to be\nswapped with the seasons. Art that\nlives and breathes.",
     "四季に合わせて掛け替えるセット。\n生きて呼吸するアート体験。"),
    ("CONCEPT 02", "Upcycled Heritage",
     "Forgotten craft materials — kimono,\ndenim, kokura-ori — reborn as canvases\nthat carry Japanese stories.",
     "忘れられた日本の素材を蘇らせ、\n物語を纏うキャンバスに。"),
    ("CONCEPT 03", "Washi as Living Paper",
     "Echizen handmade washi — colors never\nfade, fibers create organic bleeding.\nA canvas that ages beautifully.",
     "越前手漉き和紙 — 色褪せず、\n何百年も美しく経年変化するキャンバス。"),
    ("CONCEPT 04", "New Material Frontier",
     "PVC, tracing paper, translucent layers —\npushing the scroll format into contemporary\nart, photography, and graphic design.",
     "PVC・トレーシングペーパーで\n掛け軸を現代アートの領域へ拡張。"),
]

grid_x = [L_MARGIN, L_MARGIN + Inches(5.5)]
grid_y = [Inches(2.2), Inches(4.9)]
gw = Inches(5.2)
gh = Inches(2.4)

for idx, (label, title, desc_en, desc_jp) in enumerate(concepts):
    cx = grid_x[idx % 2]
    cy = grid_y[idx // 2]
    add_rect(slide, cx, cy, gw, gh, WARM_WHITE)
    add_text(slide, cx + Inches(0.3), cy + Inches(0.2), Inches(4), Inches(0.25),
             label, font_size=CAPTION_SIZE, color=LIGHT_GRAY)
    add_text(slide, cx + Inches(0.3), cy + Inches(0.5), Inches(4.5), Inches(0.4),
             title, font_size=15, color=SUMI, bold=False)
    add_text(slide, cx + Inches(0.3), cy + Inches(1.0), Inches(4.5), Inches(0.8),
             desc_en, font_size=13, color=WARM_GRAY)
    add_text(slide, cx + Inches(0.3), cy + Inches(1.75), Inches(4.5), Inches(0.5),
             desc_jp, font_size=11, color=LIGHT_GRAY)


# ================================================================
# SLIDE 11: For Artists
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_photo_placeholder(slide, R_HALF, Inches(0), HALF_W, SLIDE_H,
                      "PHOTO\nartist working on scroll /\nscroll in studio")

add_text(slide, L_MARGIN, Inches(1.0), Inches(5), Inches(0.3),
         "FOR ARTISTS", font_size=CAPTION_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(1.4), Inches(5.2), Inches(0.8),
         "A New Canvas\nfor Your Art", font_size=28, color=SUMI)

add_text(slide, L_MARGIN, Inches(2.3), Inches(5), Inches(0.4),
         "あなたのアートに、新しいキャンバスを",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

add_line(slide, L_MARGIN, Inches(2.9), Inches(0.6))

points = [
    "Paint, draw, print, photograph\n— on washi, denim, fabric, or PVC",
    "Your work becomes a rollable,\nportable, collectible piece",
    "Exhibit and sell through our platform\nand pop-up events",
]
y = 3.3
for pt in points:
    add_text(slide, L_MARGIN, Inches(y), Inches(5.2), Inches(0.6),
             pt, font_size=BODY_EN_SIZE, color=WARM_GRAY)
    y += 0.7

add_line(slide, L_MARGIN, Inches(y + 0.1), Inches(0.6))

add_text(slide, L_MARGIN, Inches(y + 0.4), Inches(5.2), Inches(0.6),
         "We provide the canvas.\nYou bring the art.",
         font_size=17, color=AI_BLUE)

add_text(slide, L_MARGIN, Inches(y + 1.05), Inches(5.2), Inches(0.4),
         "キャンバスは私たちが。アートはあなたが。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)


# ================================================================
# SLIDE 12: Partners (3 locations)
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, L_MARGIN, Inches(0.6), FULL_W, Inches(0.3),
         "PARTNERSHIP", font_size=CAPTION_SIZE, color=LIGHT_GRAY)

add_text(slide, L_MARGIN, Inches(1.0), FULL_W, Inches(0.6),
         "Partners We Seek", font_size=26, color=SUMI)

add_text(slide, L_MARGIN, Inches(1.6), FULL_W, Inches(0.3),
         "探しているパートナー", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

partners = [
    ("JAPAN", "Manufacturing & Materials",
     "Washi production (traceable, EUDR-ready)\nTextile sourcing (kimono, denim, kokura-ori)\nScroll mounting & assembly\nHardware (bamboo, reclaimed wood)",
     "和紙製造 / 生地調達 / 表装・組立 / 棒材"),
    ("EURO ZONE", "Distribution & Community",
     "Logistics & warehousing (Amsterdam hub)\nGallery & showroom partnerships\nArtist network & collaborations\nEUDR compliance & import support",
     "物流 / ギャラリー提携 / アーティスト / EUDR対応"),
    ("LONDON", "Brand & Creative",
     "Brand storytelling & content creation\nPop-up events & tea ceremony\nPhotography & studio partnerships\nUK market development",
     "ブランド発信 / イベント / 撮影 / UK市場開拓"),
]

pw = Inches(3.4)
ph = Inches(4.5)
py = Inches(2.2)

for i, (loc, subtitle, desc_en, desc_jp) in enumerate(partners):
    px = L_MARGIN + (pw + Inches(0.35)) * i
    add_rect(slide, px, py, pw, ph, WARM_WHITE)
    add_text(slide, px + Inches(0.3), py + Inches(0.25), Inches(2.8), Inches(0.25),
             loc, font_size=CAPTION_SIZE, color=AI_BLUE)
    add_text(slide, px + Inches(0.3), py + Inches(0.6), Inches(2.8), Inches(0.4),
             subtitle, font_size=16, color=SUMI, bold=False)
    add_text(slide, px + Inches(0.3), py + Inches(1.2), Inches(2.8), Inches(2.0),
             desc_en, font_size=13, color=WARM_GRAY)
    add_text(slide, px + Inches(0.3), py + Inches(3.5), Inches(2.8), Inches(0.6),
             desc_jp, font_size=10, color=LIGHT_GRAY)


# ================================================================
# SLIDE 13: Roadmap
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, L_MARGIN, Inches(0.7), FULL_W, Inches(0.6),
         "Roadmap", font_size=28, color=SUMI)

phases = [
    ("0", "Development", "NOW", "Original kakejiku R&D. Material sourcing & manufacturer partnerships.",
     "オリジナル掛け軸の研究開発。素材調達と製造パートナーの開拓。"),
    ("1", "Market Research", "", "Curated vintage sales on eBay. Studio photography in European interiors. Owned media launch.",
     "セレクト掛け軸のeBay販売、欧米インテリアでの撮影、オウンドメディア立ち上げ。"),
    ("2", "Artist Collaboration", "", "First collabs with Europe-based artists. Pop-up exhibitions in London & Amsterdam.",
     "欧州アーティストとの初コラボ。ロンドン・アムステルダムでポップアップ展示。"),
    ("3", "Scale", "", "Showroom & experience space. Subscription model. Tea ceremony events & community.",
     "ショールーム展開。サブスクリプション。お茶会イベントとコミュニティ。"),
]

y = 1.7
for num, title, badge, desc_en, desc_jp in phases:
    add_text(slide, L_MARGIN, Inches(y), Inches(0.6), Inches(0.6),
             num, font_size=40, color=SAND)
    title_text = title + ("  " + badge if badge else "")
    add_text(slide, L_MARGIN + Inches(0.8), Inches(y), Inches(9), Inches(0.35),
             title_text, font_size=16, color=SUMI, bold=False)
    add_text(slide, L_MARGIN + Inches(0.8), Inches(y + 0.38), Inches(9), Inches(0.4),
             desc_en, font_size=13, color=WARM_GRAY)
    add_text(slide, L_MARGIN + Inches(0.8), Inches(y + 0.75), Inches(9), Inches(0.3),
             desc_jp, font_size=11, color=LIGHT_GRAY)
    y += 1.35


# ================================================================
# SLIDE 14: Team
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, L_MARGIN, Inches(0.8), FULL_W, Inches(0.6),
         "Team", font_size=28, color=SUMI)

# Yayoi
add_photo_placeholder(slide, L_MARGIN, Inches(1.8), Inches(1.0), Inches(1.0), "PHOTO")
add_text(slide, L_MARGIN, Inches(3.0), Inches(4.5), Inches(0.4),
         "Yayoi Yanagisawa", font_size=20, color=SUMI)
add_text(slide, L_MARGIN, Inches(3.4), Inches(4.5), Inches(0.25),
         "Co-founder — Marketing & Creative", font_size=13, color=AI_BLUE)
add_text(slide, L_MARGIN, Inches(3.65), Inches(4.5), Inches(0.25),
         "London", font_size=12, color=LIGHT_GRAY)
add_text(slide, L_MARGIN, Inches(4.1), Inches(4.5), Inches(1.0),
         "Doshisha Univ. Commerce → Nomura Securities → Founded apparel D2C (acquired by Korean unicorn) → India jewelry D2C → London, marketing & creative",
         font_size=13, color=WARM_GRAY)
add_text(slide, L_MARGIN, Inches(5.3), Inches(4.5), Inches(0.6),
         "同志社大学商学部→野村證券→アパレルD2C（韓国ユニコーンにバイアウト）→インドD2C→ロンドン",
         font_size=11, color=LIGHT_GRAY)

# Divider
add_rect(slide, Inches(6.4), Inches(1.8), Pt(1), Inches(4.5), SAND)

# Mako
mx = Inches(7.0)
add_photo_placeholder(slide, mx, Inches(1.8), Inches(1.0), Inches(1.0), "PHOTO")
add_text(slide, mx, Inches(3.0), Inches(4.5), Inches(0.4),
         "Mako", font_size=20, color=SUMI)
add_text(slide, mx, Inches(3.4), Inches(4.5), Inches(0.25),
         "Co-founder — Finance & Operations", font_size=13, color=AI_BLUE)
add_text(slide, mx, Inches(3.65), Inches(4.5), Inches(0.25),
         "Amsterdam", font_size=12, color=LIGHT_GRAY)
add_text(slide, mx, Inches(4.1), Inches(4.5), Inches(1.0),
         "Doshisha Univ. Commerce → Certified Public Accountant → Independent consultant supporting Japanese companies expanding overseas",
         font_size=13, color=WARM_GRAY)
add_text(slide, mx, Inches(5.3), Inches(4.5), Inches(0.6),
         "同志社大学商学部→会計士→独立、日本企業の海外進出をサポート（アムステルダム拠点）",
         font_size=11, color=LIGHT_GRAY)


# ================================================================
# SLIDE 15: Vision & Close
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, SUMI)

add_text(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(0.3),
         "OUR VISION", font_size=CAPTION_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(2.2), Inches(10.3), Inches(2.0),
         "A global art platform\nwhere the world's artists create\non a Japanese format,\nand seasons return to modern homes.",
         font_size=30, color=KINARI, alignment=PP_ALIGN.CENTER)

add_line(slide, Inches(5.8), Inches(4.5), Inches(1.7))
for shape in slide.shapes:
    if hasattr(shape, 'fill') and shape.height == Pt(1):
        try:
            shape.fill.fore_color.rgb = WARM_GRAY
        except:
            pass

add_text(slide, Inches(1.5), Inches(4.9), Inches(10.3), Inches(0.8),
         "世界のアーティストが日本のフォーマットで創作し、\n現代の住空間に四季が戻るプラットフォームを。",
         font_size=SUB_JP_SIZE, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(5.9), Inches(10.3), Inches(0.6),
         "A Portable Window of Zen.",
         font_size=36, color=KINARI, alignment=PP_ALIGN.CENTER)


# ================================================================
# SLIDE 16: Contact
# ================================================================
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)

add_text(slide, Inches(1.5), Inches(2.0), Inches(10.3), Inches(0.4),
         "ZEN ART POSTER", font_size=CAPTION_SIZE, color=LIGHT_GRAY,
         alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(0.6),
         "Let's create together.", font_size=28, color=SUMI,
         alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(3.5), Inches(10.3), Inches(0.4),
         "一緒につくりましょう。", font_size=BODY_JP_SIZE, color=LIGHT_GRAY,
         alignment=PP_ALIGN.CENTER)

add_line(slide, Inches(6.0), Inches(4.2), Inches(1.3))

add_text(slide, Inches(1.5), Inches(4.7), Inches(10.3), Inches(0.4),
         "[email]", font_size=SUB_EN_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.4),
         "[website]", font_size=SUB_EN_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(5.7), Inches(10.3), Inches(0.4),
         "[Instagram]", font_size=SUB_EN_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)


# ================================================================
# TEMPLATE SLIDES
# ================================================================

# --- Template A: Title (Light) ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_text(slide, Inches(1.5), Inches(2.0), Inches(10.3), Inches(0.4),
         "SECTION LABEL", font_size=CAPTION_SIZE, color=LIGHT_GRAY,
         alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(2.7), Inches(10.3), Inches(1.0),
         "Title Goes Here", font_size=H1_SIZE, color=SUMI,
         alignment=PP_ALIGN.CENTER)
add_line(slide, Inches(6.0), Inches(3.9), Inches(1.3))
add_text(slide, Inches(1.5), Inches(4.3), Inches(10.3), Inches(0.4),
         "日本語サブタイトル", font_size=SUB_JP_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)

# --- Template B: Title (Dark) ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, SUMI)
add_text(slide, Inches(1.5), Inches(2.0), Inches(10.3), Inches(0.4),
         "SECTION LABEL", font_size=CAPTION_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(2.7), Inches(10.3), Inches(1.0),
         "Title Goes Here", font_size=H1_SIZE, color=KINARI,
         alignment=PP_ALIGN.CENTER)
add_line(slide, Inches(6.0), Inches(3.9), Inches(1.3))
add_text(slide, Inches(1.5), Inches(4.3), Inches(10.3), Inches(0.4),
         "日本語サブタイトル", font_size=SUB_JP_SIZE, color=LIGHT_GRAY,
         alignment=PP_ALIGN.CENTER)

# --- Template C: Left Text + Right Photo ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_photo_placeholder(slide, R_HALF, Inches(0), HALF_W, SLIDE_H, "PHOTO")
add_text(slide, L_MARGIN, Inches(1.5), Inches(5), Inches(0.3),
         "LABEL", font_size=CAPTION_SIZE, color=LIGHT_GRAY)
add_text(slide, L_MARGIN, Inches(2.0), Inches(5), Inches(0.7),
         "Heading Here", font_size=28, color=SUMI)
add_text(slide, L_MARGIN, Inches(2.8), Inches(5), Inches(0.4),
         "日本語見出し", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)
add_line(slide, L_MARGIN, Inches(3.4), Inches(0.6))
add_text(slide, L_MARGIN, Inches(3.8), Inches(5), Inches(0.8),
         "Body text goes here. Keep it concise\nand let the photo do the talking.",
         font_size=BODY_EN_SIZE, color=WARM_GRAY)
add_text(slide, L_MARGIN, Inches(4.7), Inches(5), Inches(0.6),
         "本文テキスト。簡潔に。\n写真に語らせる。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

# --- Template D: Left Photo + Right Text ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_photo_placeholder(slide, Inches(0), Inches(0), HALF_W, SLIDE_H, "PHOTO")
add_text(slide, Inches(7.2), Inches(1.5), Inches(5), Inches(0.3),
         "LABEL", font_size=CAPTION_SIZE, color=LIGHT_GRAY)
add_text(slide, Inches(7.2), Inches(2.0), Inches(5), Inches(0.7),
         "Heading Here", font_size=28, color=SUMI)
add_text(slide, Inches(7.2), Inches(2.8), Inches(5), Inches(0.4),
         "日本語見出し", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)
add_line(slide, Inches(7.2), Inches(3.4), Inches(0.6))
add_text(slide, Inches(7.2), Inches(3.8), Inches(5), Inches(0.8),
         "Body text goes here. Keep it concise\nand let the photo do the talking.",
         font_size=BODY_EN_SIZE, color=WARM_GRAY)
add_text(slide, Inches(7.2), Inches(4.7), Inches(5), Inches(0.6),
         "本文テキスト。簡潔に。\n写真に語らせる。",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

# --- Template E: Quote (Light) ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_text(slide, Inches(1.5), Inches(2.5), Inches(10.3), Inches(1.0),
         '"Quote or key statement goes here.\nKeep it to one or two lines."',
         font_size=22, color=SUMI, alignment=PP_ALIGN.CENTER)
add_line(slide, Inches(6.0), Inches(3.9), Inches(1.3))
add_text(slide, Inches(1.5), Inches(4.2), Inches(10.3), Inches(0.8),
         "「引用文やキーステートメント。\n1〜2行に収める。」",
         font_size=16, color=WARM_GRAY, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.3),
         "— ATTRIBUTION", font_size=CAPTION_SIZE, color=LIGHT_GRAY,
         alignment=PP_ALIGN.CENTER)

# --- Template F: Quote (Dark) ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, SUMI)
add_text(slide, Inches(1.5), Inches(2.5), Inches(10.3), Inches(1.0),
         '"Quote or key statement goes here.\nKeep it to one or two lines."',
         font_size=22, color=KINARI, alignment=PP_ALIGN.CENTER)
add_line(slide, Inches(6.0), Inches(3.9), Inches(1.3))
add_text(slide, Inches(1.5), Inches(4.2), Inches(10.3), Inches(0.8),
         "「引用文やキーステートメント。\n1〜2行に収める。」",
         font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.3),
         "— ATTRIBUTION", font_size=CAPTION_SIZE, color=WARM_GRAY,
         alignment=PP_ALIGN.CENTER)

# --- Template G: 2-Column Cards ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_text(slide, L_MARGIN, Inches(0.6), FULL_W, Inches(0.3),
         "LABEL", font_size=CAPTION_SIZE, color=LIGHT_GRAY)
add_text(slide, L_MARGIN, Inches(1.0), FULL_W, Inches(0.6),
         "Section Heading", font_size=28, color=SUMI)
add_text(slide, L_MARGIN, Inches(1.6), FULL_W, Inches(0.3),
         "セクション見出し", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)
for i in range(2):
    cx = L_MARGIN + Inches(5.5) * i
    add_rect(slide, cx, Inches(2.3), Inches(5.2), Inches(4.5), WARM_WHITE)
    add_text(slide, cx + Inches(0.35), Inches(2.6), Inches(4.5), Inches(0.25),
             f"CARD 0{i+1}", font_size=CAPTION_SIZE, color=AI_BLUE)
    add_text(slide, cx + Inches(0.35), Inches(3.0), Inches(4.5), Inches(0.4),
             "Card Title", font_size=18, color=SUMI)
    add_text(slide, cx + Inches(0.35), Inches(3.6), Inches(4.5), Inches(0.8),
             "Description text here.\nKeep it to 2-3 lines.",
             font_size=BODY_EN_SIZE, color=WARM_GRAY)
    add_text(slide, cx + Inches(0.35), Inches(4.5), Inches(4.5), Inches(0.5),
             "説明テキスト。\n2〜3行に収める。",
             font_size=BODY_JP_SIZE, color=LIGHT_GRAY)

# --- Template H: 3-Column Cards ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_text(slide, L_MARGIN, Inches(0.6), FULL_W, Inches(0.3),
         "LABEL", font_size=CAPTION_SIZE, color=LIGHT_GRAY)
add_text(slide, L_MARGIN, Inches(1.0), FULL_W, Inches(0.6),
         "Section Heading", font_size=28, color=SUMI)
add_text(slide, L_MARGIN, Inches(1.6), FULL_W, Inches(0.3),
         "セクション見出し", font_size=BODY_JP_SIZE, color=LIGHT_GRAY)
for i in range(3):
    cx = L_MARGIN + (Inches(3.4) + Inches(0.35)) * i
    add_rect(slide, cx, Inches(2.3), Inches(3.4), Inches(4.5), WARM_WHITE)
    add_text(slide, cx + Inches(0.3), Inches(2.6), Inches(2.8), Inches(0.25),
             f"0{i+1}", font_size=CAPTION_SIZE, color=LIGHT_GRAY)
    add_text(slide, cx + Inches(0.3), Inches(3.0), Inches(2.8), Inches(0.4),
             "Card Title", font_size=16, color=SUMI)
    add_text(slide, cx + Inches(0.3), Inches(3.5), Inches(2.8), Inches(0.6),
             "Description text.",
             font_size=13, color=WARM_GRAY)
    add_text(slide, cx + Inches(0.3), Inches(4.2), Inches(2.8), Inches(0.4),
             "説明テキスト。",
             font_size=11, color=LIGHT_GRAY)

# --- Template I: Large Number ---
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, KINARI)
add_text(slide, Inches(1.5), Inches(1.8), Inches(10.3), Inches(1.5),
         "87%", font_size=120, color=SAND, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(3.8), Inches(10.3), Inches(0.6),
         "Statement about the number",
         font_size=H3_SIZE, color=SUMI, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(4.5), Inches(10.3), Inches(0.4),
         "数字についての説明文",
         font_size=BODY_JP_SIZE, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(5.3), Inches(10.3), Inches(0.3),
         "SOURCE: Attribution here",
         font_size=CAPTION_SIZE, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# ================================================================
# SAVE
# ================================================================
output_path = '/Users/yayoi/clawd/zen-art-poster/zen-art-poster.pptx'
prs.save(output_path)
print(f"Saved to {output_path}")
print(f"Total slides: {len(prs.slides)}")
